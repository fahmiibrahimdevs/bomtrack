import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any

def generate_bom_excel(project_name: str, items: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM Summary"
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # Title Row
    ws.merge_cells("A1:I1")
    ws["A1"] = f"BILL OF MATERIALS (BOM) — {project_name.upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 35
    
    # Headers
    headers = [
        "No", "Nama Komponen / Barang", "Varian", "Kategori", 
        "Toko", "Lokasi", "Qty", "Harga Satuan (Rp)", "Total Harga (Rp)", "Link Pembelian", "Status"
    ]
    
    ws.row_dimensions[3].height = 25
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    # Rows
    current_row = 4
    total_cost = 0.0
    for idx, item in enumerate(items, 1):
        unit_price = float(item.get("unit_price", 0))
        qty = int(item.get("quantity", 1))
        row_total = float(item.get("total_price", unit_price * qty))
        total_cost += row_total
        
        status_text = "Sudah Dibeli" if item.get("is_purchased") == 1 else "Rencana"
        
        row_data = [
            idx,
            item.get("name", ""),
            item.get("variant_name", "-"),
            item.get("category", "General"),
            item.get("shop_name", "-"),
            item.get("shop_location", "-"),
            qty,
            unit_price,
            row_total,
            item.get("product_url", ""),
            status_text
        ]
        
        ws.row_dimensions[current_row].height = 22
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = val
            cell.border = thin_border
            
            # Format numbers
            if col_num == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num in (7,):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "#,##0"
            elif col_num in (8, 9):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "Rp #,##0"
            elif col_num == 10 and val:
                cell.hyperlink = val
                cell.font = Font(color="2563EB", underline="single")
            elif col_num == 11:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
        current_row += 1
        
    # Total Summary Row
    ws.row_dimensions[current_row].height = 25
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
    summary_cell = ws.cell(row=current_row, column=1)
    summary_cell.value = "TOTAL ESTIMASI BIAYA BOM:"
    summary_cell.font = bold_font
    summary_cell.alignment = Alignment(horizontal="right", vertical="center")
    summary_cell.fill = total_fill
    
    for c in range(1, 9):
        ws.cell(row=current_row, column=c).border = thin_border
        ws.cell(row=current_row, column=c).fill = total_fill
        
    total_val_cell = ws.cell(row=current_row, column=9)
    total_val_cell.value = total_cost
    total_val_cell.font = bold_font
    total_val_cell.number_format = "Rp #,##0"
    total_val_cell.alignment = Alignment(horizontal="right", vertical="center")
    total_val_cell.fill = total_fill
    total_val_cell.border = thin_border
    
    ws.cell(row=current_row, column=10).border = thin_border
    ws.cell(row=current_row, column=10).fill = total_fill
    ws.cell(row=current_row, column=11).border = thin_border
    ws.cell(row=current_row, column=11).fill = total_fill

    # Column Widths
    col_widths = {
        "A": 6,   # No
        "B": 35,  # Nama
        "C": 20,  # Varian
        "D": 15,  # Kategori
        "E": 22,  # Toko
        "F": 18,  # Lokasi
        "G": 8,   # Qty
        "H": 18,  # Harga Satuan
        "I": 20,  # Total Harga
        "J": 30,  # Link
        "K": 15,  # Status
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
        
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
